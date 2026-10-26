"""Stage article topics from Keyword Planner exports (see references/topic-research.md).

    python3 topic_rank.py topics.json topics.xlsx export1.csv [export2.csv ...]
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from location_rank import read_export

# Monthly searches in the country. A topic proven on a sister site needs only measurable volume.
LAUNCH, WAVE = 100, 30


def stage(vol, proven):
    if vol >= LAUNCH or (proven and vol >= 10):
        return 'Launch'
    if vol >= WAVE or proven:
        return 'Wave 2'
    return 'Hold'


def check_against(kw, money):
    # the money page whose search shares the topic's head noun: run the top-10 overlap test against it
    heads = {'fascia': 'fascias and soffits', 'soffit': 'fascias and soffits', 'roofline': 'fascias and soffits', 'downpipe': 'downpipe repairs',
             'replace': 'gutter replacement', 'repair': 'gutter repairs', 'guttering': 'guttering', 'gutter': 'guttering'}  # most specific first
    for h, search in heads.items():
        if h in kw:
            return next((f'{u} ("{q}")' for u, q in money.items() if q.startswith(search)), '')
    return ''


if __name__ == '__main__':
    cfg = json.load(open(sys.argv[1]))
    vol = {}
    for p in sys.argv[3:]: vol.update(read_export(p, cfg['country'], with_months=True))
    proven, money = cfg.get('proven', {}), cfg.get('money_pages', {})
    rows = []
    for g, kws in cfg['groups'].items():
        for k in kws:
            v, months = vol.get(k.lower(), (0, ()))
            rows.append(dict(group=g, kw=k, vol=v, months=months, proven=proven.get(k, ''), stage=stage(v, k in proven)))
    # identical volume and identical monthly pattern inside a group: Google treats them as one search
    for r in rows:
        twin = next((o['kw'] for o in rows if o is not r and o['group'] == r['group'] and r['vol'] >= 10 and o['vol'] == r['vol'] and o['months'] == r['months']), '')
        r['twin'] = twin
    wb = Workbook(); ws = wb.active; ws.title = 'Topics'
    ws.append(['Group', 'Topic search', 'Searches / month', 'Stage', 'Sister-site evidence', 'Probably one search with', 'Run the top-10 overlap test against', 'Overlap % (fill in)', 'Decision (fill in)'])
    fill = {'Launch': 'E2F5E6', 'Wave 2': 'FFF4D6'}
    for r in sorted(rows, key=lambda r: (list(cfg['groups']).index(r['group']), -r['vol'])):
        ws.append([r['group'], r['kw'], r['vol'], r['stage'], r['proven'], r['twin'], check_against(r['kw'], money) if r['group'].startswith('Cost') or r['group'] == 'Problem and signs' else '', '', ''])
        if r['stage'] in fill:
            for x in ws[ws.max_row][:4]: x.fill = PatternFill('solid', fgColor=fill[r['stage']])
    for i, w in enumerate([24, 40, 12, 10, 60, 30, 40, 12, 30], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1A1A1A'); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
    m = wb.create_sheet('Method')
    for l in [f'Keyword Planner, {cfg["country"]}, average monthly searches. Stage: Launch at {LAUNCH}+ or proven on a sister site with measurable volume; Wave 2 at {WAVE}+ or proven; otherwise Hold.',
              'Before a topic is written: run the top-10 overlap test against the money page named. 60% or more of the same hostnames = one search: it becomes a section of that page, not an article.',
              'Topics with the same volume and the same monthly pattern are probably one search to Google: one article covers both.',
              'Cost in a place: a spoke is written only where the place version has its own volume or the sister site shows place spokes earning.',
              'Each site writes its own article. Topics are shared between sites; text never is.']:
        m.append([l])
    m.column_dimensions['A'].width = 150
    wb.save(sys.argv[2])
    from collections import Counter
    print(dict(Counter(r['stage'] for r in rows)), 'twins flagged:', sum(1 for r in rows if r['twin']))
